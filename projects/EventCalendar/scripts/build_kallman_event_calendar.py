from __future__ import annotations

import json
import re
from collections import Counter
from datetime import date, datetime, time
from pathlib import Path

import openpyxl


ROOT_DIR = Path(__file__).resolve().parents[1]
SOURCE_WORKBOOK = Path(
    r"C:\Users\kylep\Kallman Worldwide, Inc\Data Warehouse - Documents\TASS PQ.xlsx"
)
STATUS_WORKBOOK = Path(
    r"C:\Users\kylep\Kallman Worldwide, Inc\Data Warehouse - Documents\Event Status.xlsx"
)
OUTPUT_PATH = ROOT_DIR / "outputs" / "kallman_event_calendar.html"
GENERATED_ON = "2026-06-10"

PREFERRED_CATEGORY_COLORS = {
    "Aerospace": "#0B4B93",
    "Defense & Security": "#C92127",
    "Health & Medical": "#0F766E",
    "Energy": "#D97706",
    "Maritime": "#2D6CB9",
    "Mining": "#6D5BD0",
    "Pets": "#BE185D",
    "Fitness": "#16A34A",
    "Agriculture & Food": "#7C3AED",
    "Other": "#475569",
    "Uncategorized": "#64748B",
}

FALLBACK_PALETTE = [
    "#0B4B93",
    "#C92127",
    "#0F766E",
    "#D97706",
    "#6D5BD0",
    "#BE185D",
    "#2D6CB9",
    "#16A34A",
    "#7C3AED",
    "#475569",
    "#0891B2",
    "#A16207",
]


def clean(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def normalized(value) -> str:
    return clean(value).casefold()


def as_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if not value:
        return None
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%Y %H:%M"):
        try:
            return datetime.strptime(str(value), fmt).date()
        except ValueError:
            pass
    return None


def as_time(value) -> time | None:
    if isinstance(value, datetime):
        return value.time().replace(second=0, microsecond=0)
    if isinstance(value, time):
        return value.replace(second=0, microsecond=0)
    if not value:
        return None
    text = str(value).strip()
    for fmt in ("%H:%M:%S", "%H:%M", "%I:%M %p", "%m/%d/%Y %H:%M"):
        try:
            return datetime.strptime(text, fmt).time().replace(second=0, microsecond=0)
        except ValueError:
            pass
    return None


def time_to_iso(value) -> str:
    parsed = as_time(value)
    return parsed.strftime("%H:%M") if parsed else ""


def ole_color_to_hex(value) -> str:
    try:
        number = int(value)
    except (TypeError, ValueError):
        return ""
    red = number & 0xFF
    green = (number >> 8) & 0xFF
    blue = (number >> 16) & 0xFF
    return f"#{red:02X}{green:02X}{blue:02X}"


def read_sheet_rows(workbook, sheet_name: str):
    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(values_only=True)
    headers = [clean(value) for value in next(rows)]
    indexes = {header: index for index, header in enumerate(headers) if header}
    return headers, indexes, rows


def read_lookup(workbook, sheet_name: str) -> dict[str, str]:
    if sheet_name not in workbook.sheetnames:
        return {}
    _, indexes, rows = read_sheet_rows(workbook, sheet_name)
    desc_index = indexes.get("Description")
    code_index = indexes.get("Code")
    if desc_index is None or code_index is None:
        return {}
    lookup = {}
    for row in rows:
        description = clean(row[desc_index] if desc_index < len(row) else "")
        code = clean(row[code_index] if code_index < len(row) else "")
        if description and code and description != "Description":
            lookup[code] = description
    return lookup


def read_status_colors() -> tuple[dict[str, dict[str, str]], dict[str, str]]:
    workbook = openpyxl.load_workbook(STATUS_WORKBOOK, read_only=True, data_only=True)
    sheet = workbook.active
    headers = [clean(cell.value) for cell in sheet[1]]
    indexes = {header: index for index, header in enumerate(headers) if header}
    by_status: dict[str, dict[str, str]] = {}
    descriptions_by_code: dict[str, str] = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        description = clean(row[indexes["Description"]])
        code = clean(row[indexes["Code"]])
        background = ole_color_to_hex(row[indexes["Background Color"]])
        text = ole_color_to_hex(row[indexes["Text Color"]]) or "#FFFFFF"
        if description:
            by_status[description] = {"background": background, "text": text, "code": code}
        if code and description:
            descriptions_by_code[code] = description
    return by_status, descriptions_by_code


def read_category_hexes(workbook) -> dict[str, str]:
    candidates = [
        "Categories",
        "Category Selection",
        "CategorySelection",
        "ClassesEvents",
        "Event Classes",
    ]
    color_map: dict[str, str] = {}
    for sheet_name in candidates:
        if sheet_name not in workbook.sheetnames:
            continue
        sheet = workbook[sheet_name]
        headers = [clean(cell.value).lower() for cell in sheet[1]]
        category_index = next(
            (
                i
                for i, header in enumerate(headers)
                if header in {"category", "event class", "class", "description"}
            ),
            None,
        )
        color_index = next(
            (
                i
                for i, header in enumerate(headers)
                if "hex" in header or ("color" in header and "text" not in header)
            ),
            None,
        )
        if category_index is None or color_index is None:
            continue
        for row in sheet.iter_rows(min_row=2, values_only=True):
            category = clean(row[category_index] if category_index < len(row) else "")
            color = clean(row[color_index] if color_index < len(row) else "").upper()
            if category and re.fullmatch(r"#?[0-9A-F]{6}", color):
                color_map[category] = color if color.startswith("#") else f"#{color}"
    return color_map


def build_events() -> tuple[list[dict], dict[str, str], dict[str, dict[str, str]], dict]:
    workbook = openpyxl.load_workbook(SOURCE_WORKBOOK, read_only=True, data_only=True)
    class_lookup = read_lookup(workbook, "ClassesEvents")
    type_lookup = read_lookup(workbook, "TypeEvents")
    status_colors, statuses_by_code = read_status_colors()

    _, raw_indexes, raw_rows = read_sheet_rows(workbook, "Events_Pull")
    raw_by_key = {}
    for row in raw_rows:
        start = as_date(row[raw_indexes["StartDate"]])
        description = clean(row[raw_indexes["Description"]])
        if not start or not description:
            continue
        key = (normalized(description), start.isoformat())
        raw_by_key[key] = {
            "eventId": row[raw_indexes["EventID"]],
            "startTime": time_to_iso(row[raw_indexes["StartTime"]]),
            "endTime": time_to_iso(row[raw_indexes["EndTime"]]),
            "statusCode": clean(row[raw_indexes["Status"]]),
            "classCode": clean(row[raw_indexes["Class"]]),
            "typeCode": clean(row[raw_indexes["Type"]]),
        }

    _, indexes, rows = read_sheet_rows(workbook, "TASS (new)")
    events: list[dict] = []
    for row in rows:
        title = clean(row[indexes["Description"]])
        start = as_date(row[indexes["Start Date"]])
        end = as_date(row[indexes["End Date"]]) or start
        if not title or not start:
            continue

        raw = raw_by_key.get((normalized(title), start.isoformat()), {})
        status_code = raw.get("statusCode", "")
        status = clean(row[indexes["Event Status"]]) or statuses_by_code.get(status_code, "")
        category_code = raw.get("classCode", "")
        category = clean(row[indexes["Event Class"]]) or class_lookup.get(category_code, "Uncategorized")
        event_type = clean(row[indexes["Event Type"]]) or type_lookup.get(raw.get("typeCode", ""), "")
        city = clean(row[indexes["City"]])
        country = clean(row[indexes["Country"]])
        location = ", ".join(part for part in (city, country) if part) or "Location not listed"
        updated = as_date(row[indexes["Event Details Last Updated"]])

        team_parts = []
        for role in ("PM", "Sales", "PC", "OPS", "CS", "Exhibitor Services"):
            value = clean(row[indexes[role]])
            if value:
                team_parts.append(f"{role}: {value}")

        events.append(
            {
                "id": len(events) + 1,
                "eventId": clean(raw.get("eventId", "")),
                "title": title,
                "startDate": start.isoformat(),
                "endDate": end.isoformat() if end else start.isoformat(),
                "startTime": raw.get("startTime", "08:00"),
                "endTime": raw.get("endTime", "17:00"),
                "location": location,
                "city": city,
                "country": country,
                "category": category or "Uncategorized",
                "categoryCode": category_code,
                "type": event_type or "Unspecified",
                "typeCode": raw.get("typeCode", ""),
                "status": status or "Unspecified",
                "statusCode": status_code,
                "foundation": clean(row[indexes["Kallman Foundation Event"]]) or "No",
                "team": "; ".join(team_parts),
                "updated": updated.isoformat() if updated else "",
            }
        )

    workbook_hexes = read_category_hexes(workbook)
    workbook.close()
    category_colors = assign_category_colors(events, workbook_hexes)
    for event in events:
        event["categoryColor"] = category_colors[event["category"]]
        status_style = status_colors.get(event["status"], {})
        event["statusBackground"] = status_style.get("background", "#64748B")
        event["statusText"] = status_style.get("text", "#FFFFFF")

    metadata = {
        "sourceWorkbook": str(SOURCE_WORKBOOK),
        "statusWorkbook": str(STATUS_WORKBOOK),
        "generatedOn": GENERATED_ON,
        "eventCount": len(events),
        "dateStart": min(event["startDate"] for event in events),
        "dateEnd": max(event["endDate"] for event in events),
        "categoryCount": len({event["category"] for event in events}),
        "locationCount": len({event["location"] for event in events}),
        "defaultMonth": default_month(events),
        "categoriesByCount": Counter(event["category"] for event in events),
    }
    return events, category_colors, status_colors, metadata


def assign_category_colors(events: list[dict], workbook_hexes: dict[str, str]) -> dict[str, str]:
    categories = sorted({event["category"] for event in events})
    category_colors: dict[str, str] = {}
    for index, category in enumerate(categories):
        category_colors[category] = (
            workbook_hexes.get(category)
            or PREFERRED_CATEGORY_COLORS.get(category)
            or FALLBACK_PALETTE[index % len(FALLBACK_PALETTE)]
        )
    return category_colors


def default_month(events: list[dict]) -> str:
    today = datetime.strptime(GENERATED_ON, "%Y-%m-%d").date()
    upcoming = sorted(event["startDate"] for event in events if event["endDate"] >= today.isoformat())
    if upcoming:
        return upcoming[0][:7]
    return min(event["startDate"] for event in events)[:7]


HTML_TEMPLATE = r"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <title>Kallman Event Calendar</title>
  <style>
    :root {
      --kallman-blue: #082f6f;
      --kallman-blue-2: #0b4b93;
      --kallman-red: #c92127;
      --kallman-gold: #f3b21a;
      --ink: #172033;
      --muted: #64748b;
      --line: #d9e1ea;
      --line-soft: #eef2f7;
      --bg: #f5f7fb;
      --panel: #ffffff;
      --shadow: 0 12px 30px rgba(8,47,111,.10);
    }
    * { box-sizing: border-box; }
    body { margin: 0; font-family: Arial, Helvetica, sans-serif; background: var(--bg); color: var(--ink); }
    button, input, select { font: inherit; }
    .topbar {
      color: white;
      padding: 22px 28px 34px;
      border-bottom: 4px solid var(--kallman-red);
      background:
        linear-gradient(115deg, rgba(8,47,111,.98), rgba(11,75,147,.92)),
        url("https://www.kallman.com/wp-content/uploads/2023/06/52718043124_8bd4f20714_o-scaled.jpg");
      background-size: cover;
      background-position: center;
    }
    .brand { display: flex; justify-content: space-between; align-items: center; gap: 16px; max-width: 1440px; margin: 0 auto; }
    .logo { font-weight: 800; letter-spacing: .04em; text-transform: uppercase; font-size: 21px; }
    .logo span { color: var(--kallman-red); }
    .source { color: rgba(255,255,255,.78); font-size: 13px; text-align: right; }
    .hero { max-width: 1440px; margin: 20px auto 0; }
    .hero h1 { margin: 0 0 8px; font-size: clamp(30px, 4vw, 50px); line-height: 1.02; letter-spacing: 0; }
    .hero p { margin: 0; max-width: 860px; color: rgba(255,255,255,.86); font-size: 15px; line-height: 1.45; }
    .wrap { max-width: 1440px; margin: -18px auto 44px; padding: 0 20px; }
    .filters {
      display: grid;
      grid-template-columns: minmax(220px, 1.3fr) minmax(160px, .9fr) repeat(4, minmax(130px, .7fr)) minmax(180px, .9fr) auto;
      gap: 12px;
      align-items: end;
      background: var(--panel);
      border: 1px solid var(--line);
      border-radius: 8px;
      box-shadow: var(--shadow);
      padding: 16px;
      position: sticky;
      top: 10px;
      z-index: 5;
    }
    label { display: block; font-size: 11px; font-weight: 800; letter-spacing: .06em; text-transform: uppercase; color: var(--kallman-blue); margin-bottom: 6px; }
    input, select {
      width: 100%;
      border: 1px solid #cbd5e1;
      border-radius: 8px;
      padding: 10px 11px;
      font-size: 14px;
      background: white;
      min-height: 42px;
      color: var(--ink);
    }
    .reset-btn {
      border: 0;
      border-radius: 8px;
      padding: 11px 15px;
      min-height: 42px;
      background: var(--kallman-red);
      color: white;
      font-weight: 800;
      cursor: pointer;
      box-shadow: 0 8px 18px rgba(201,33,39,.22);
    }
    .summary { display: grid; grid-template-columns: repeat(4, minmax(160px, 1fr)); gap: 12px; margin: 16px 0; }
    .stat { background: white; border: 1px solid var(--line); border-radius: 8px; padding: 13px 15px; min-height: 76px; }
    .stat strong { display: block; color: var(--kallman-blue); font-size: 24px; line-height: 1.1; }
    .stat span { display: block; margin-top: 5px; color: var(--muted); font-size: 11px; text-transform: uppercase; letter-spacing: .06em; font-weight: 800; }
    .section-title { display: flex; justify-content: space-between; align-items: end; margin: 24px 0 10px; gap: 16px; }
    .section-title h2 { margin: 0; color: var(--kallman-blue); font-size: 23px; }
    .section-title small { color: var(--muted); text-align: right; line-height: 1.35; }
    .legend { display: flex; flex-wrap: wrap; gap: 8px; margin: 0 0 14px; }
    .legend-item { display: inline-flex; align-items: center; gap: 7px; padding: 7px 9px; border: 1px solid var(--line); border-radius: 999px; background: white; font-size: 12px; font-weight: 800; color: var(--kallman-blue); }
    .swatch { width: 11px; height: 11px; border-radius: 3px; background: var(--swatch); box-shadow: inset 0 0 0 1px rgba(0,0,0,.08); }
    .calendar-shell, .schedule {
      background: white;
      border: 1px solid var(--line);
      border-radius: 8px;
      overflow: hidden;
      box-shadow: 0 8px 22px rgba(15,23,42,.05);
    }
    .calendar-toolbar {
      display: flex;
      justify-content: space-between;
      align-items: center;
      gap: 12px;
      padding: 13px 14px;
      border-bottom: 1px solid var(--line);
      background: #fbfdff;
    }
    .month-controls { display: flex; align-items: center; gap: 8px; }
    .icon-btn {
      width: 38px;
      height: 38px;
      border: 1px solid #cbd5e1;
      border-radius: 8px;
      background: white;
      color: var(--kallman-blue);
      font-weight: 900;
      cursor: pointer;
    }
    .month-label { color: var(--kallman-blue); font-size: 18px; font-weight: 800; min-width: 165px; }
    .month-select { max-width: 190px; min-height: 38px; padding: 8px 10px; }
    .calendar-count { color: var(--muted); font-size: 13px; font-weight: 700; }
    .month-grid { display: grid; grid-template-columns: repeat(7, minmax(0, 1fr)); }
    .month-dow { padding: 10px; text-align: center; background: var(--kallman-blue); color: white; border-right: 1px solid rgba(255,255,255,.14); font-size: 11px; font-weight: 800; text-transform: uppercase; letter-spacing: .06em; }
    .month-cell { min-height: 148px; border-top: 1px solid var(--line); border-right: 1px solid var(--line-soft); padding: 8px; background: white; }
    .month-cell:nth-child(7n) { border-right: 0; }
    .month-cell.outside { background: #f8fafc; color: #98a2b3; }
    .month-date { display: flex; justify-content: space-between; align-items: center; color: var(--kallman-blue); font-size: 13px; font-weight: 800; margin-bottom: 6px; }
    .outside .month-date { color: #98a2b3; }
    .today .month-date span:first-child { display: inline-flex; align-items: center; justify-content: center; min-width: 24px; height: 24px; border-radius: 999px; background: var(--kallman-red); color: white; }
    .day-total { color: var(--muted); font-size: 11px; font-weight: 700; }
    .calendar-event {
      display: block;
      width: 100%;
      border: 1px solid rgba(15,23,42,.08);
      border-left: 4px solid var(--event-color);
      border-radius: 6px;
      background: var(--event-bg);
      color: var(--ink);
      padding: 6px 7px;
      margin: 5px 0 0;
      text-align: left;
      cursor: pointer;
      overflow: hidden;
    }
    .calendar-event:hover, .schedule-row:hover { box-shadow: inset 0 0 0 999px rgba(8,47,111,.035); }
    .calendar-event strong { display: block; font-size: 11px; line-height: 1.25; white-space: nowrap; overflow: hidden; text-overflow: ellipsis; }
    .calendar-event span { display: block; margin-top: 2px; color: #475569; font-size: 10px; font-weight: 800; }
    .month-more { margin-top: 5px; font-size: 11px; color: var(--muted); font-weight: 700; }
    .empty { color: var(--muted); padding: 18px; text-align: center; border: 1px dashed var(--line); background: #fff; }
    .schedule-row {
      display: grid;
      grid-template-columns: 150px minmax(240px, 1fr) 190px 170px;
      gap: 14px;
      align-items: start;
      border-top: 1px solid var(--line);
      border-left: 5px solid var(--event-color);
      padding: 14px 16px;
      background: white;
      cursor: pointer;
    }
    .schedule-row:first-child { border-top: 0; }
    .date-pill { display: inline-block; background: var(--kallman-blue); color: white; border-radius: 999px; padding: 6px 10px; font-size: 12px; font-weight: 800; }
    .timeblock { color: var(--kallman-red); font-weight: 800; margin-top: 7px; font-size: 13px; }
    .titleline { font-weight: 800; color: var(--ink); margin-bottom: 5px; line-height: 1.3; }
    .details { color: var(--muted); font-size: 13px; line-height: 1.45; }
    .chipline { display: flex; flex-wrap: wrap; gap: 6px; margin-top: 8px; }
    .chip { display: inline-flex; align-items: center; border-radius: 999px; padding: 4px 8px; font-size: 11px; font-weight: 800; background: #eef2ff; color: var(--kallman-blue); }
    .status-chip { background: var(--status-bg); color: var(--status-fg); }
    .category-pill { display: inline-flex; align-items: center; gap: 7px; font-size: 13px; font-weight: 800; color: var(--kallman-blue); }
    .location-text { color: var(--ink); font-size: 13px; font-weight: 700; line-height: 1.35; }
    .modal-backdrop { display: none; position: fixed; inset: 0; background: rgba(15,23,42,.50); z-index: 20; padding: 22px; align-items: center; justify-content: center; }
    .modal-backdrop.open { display: flex; }
    .modal { width: min(760px, 100%); max-height: min(760px, calc(100vh - 44px)); overflow: auto; background: white; border-radius: 8px; box-shadow: 0 24px 70px rgba(0,0,0,.30); }
    .modal-head { display: flex; justify-content: space-between; gap: 18px; padding: 20px 22px; border-bottom: 1px solid var(--line); }
    .modal-head h3 { margin: 0 0 8px; color: var(--kallman-blue); font-size: 23px; line-height: 1.2; }
    .modal-subtitle { color: var(--muted); font-size: 13px; }
    .modal-close { border: 0; background: #eef2ff; color: var(--kallman-blue); border-radius: 8px; width: 38px; height: 38px; font-size: 24px; line-height: 1; cursor: pointer; flex: 0 0 38px; }
    .modal-body { padding: 18px 22px 22px; display: grid; grid-template-columns: repeat(2, minmax(0, 1fr)); gap: 14px; }
    .modal-field { border: 1px solid var(--line); border-radius: 8px; padding: 12px; background: #fbfdff; }
    .modal-field.wide { grid-column: 1 / -1; }
    .modal-label { font-size: 11px; font-weight: 800; text-transform: uppercase; letter-spacing: .06em; color: var(--muted); margin-bottom: 6px; }
    .modal-value { font-size: 14px; line-height: 1.45; color: var(--ink); word-break: break-word; }
    @media (max-width: 1260px) {
      .filters { grid-template-columns: repeat(4, minmax(0, 1fr)); }
      .reset-btn { width: 100%; }
    }
    @media (max-width: 900px) {
      .topbar { padding: 18px 18px 30px; }
      .wrap { padding: 0 14px; }
      .filters, .summary { grid-template-columns: 1fr 1fr; }
      .calendar-toolbar, .section-title, .brand { align-items: flex-start; flex-direction: column; }
      .source { text-align: left; }
      .month-grid { grid-template-columns: 1fr; }
      .month-dow { display: none; }
      .month-cell { min-height: auto; border-right: 0; }
      .schedule-row { grid-template-columns: 1fr; gap: 8px; }
      .modal-body { grid-template-columns: 1fr; }
    }
    @media (max-width: 620px) {
      .filters, .summary { grid-template-columns: 1fr; }
      .month-controls { flex-wrap: wrap; }
    }
  </style>
</head>
<body>
  <header class="topbar">
    <div class="brand">
      <div class="logo">Kallman <span>Worldwide</span></div>
      <div class="source">Generated __GENERATED_ON__ from TASS PQ.xlsx</div>
    </div>
    <div class="hero">
      <h1>Event Calendar</h1>
      <p>Filterable calendar and run schedule built from the workbook event columns, with event colors tied to category.</p>
    </div>
  </header>

  <main class="wrap">
    <section class="filters" aria-label="Calendar filters">
      <div>
        <label for="nameFilter">Event name</label>
        <input id="nameFilter" type="search" placeholder="Search event, team, type..." />
      </div>
      <div>
        <label for="locationFilter">Location</label>
        <select id="locationFilter"></select>
      </div>
      <div>
        <label for="dateStart">Date from</label>
        <input id="dateStart" type="date" />
      </div>
      <div>
        <label for="dateEnd">Date to</label>
        <input id="dateEnd" type="date" />
      </div>
      <div>
        <label for="timeStart">Time from</label>
        <input id="timeStart" type="time" />
      </div>
      <div>
        <label for="timeEnd">Time to</label>
        <input id="timeEnd" type="time" />
      </div>
      <div>
        <label for="categoryFilter">Category</label>
        <select id="categoryFilter"></select>
      </div>
      <button class="reset-btn" id="resetBtn" type="button">Reset</button>
    </section>

    <section class="summary" aria-label="Calendar summary">
      <div class="stat"><strong id="visibleCount">0</strong><span>Visible events</span></div>
      <div class="stat"><strong id="dateSpan">-</strong><span>Date range</span></div>
      <div class="stat"><strong id="categoryCount">0</strong><span>Categories</span></div>
      <div class="stat"><strong id="locationCount">0</strong><span>Locations</span></div>
    </section>

    <div class="section-title">
      <h2>Calendar View</h2>
      <small id="calendarHelp">Month view respects the active filters.</small>
    </div>
    <section class="legend" id="legend"></section>
    <section class="calendar-shell" aria-label="Month calendar">
      <div class="calendar-toolbar">
        <div class="month-controls">
          <button class="icon-btn" id="prevMonth" type="button" aria-label="Previous month">&lt;</button>
          <div class="month-label" id="monthLabel"></div>
          <button class="icon-btn" id="nextMonth" type="button" aria-label="Next month">&gt;</button>
          <select class="month-select" id="monthSelect" aria-label="Jump to month"></select>
        </div>
        <div class="calendar-count" id="monthCount"></div>
      </div>
      <div class="month-grid" id="calendarGrid"></div>
    </section>

    <div class="section-title">
      <h2>Run Schedule</h2>
      <small id="scheduleHelp">Sorted by start date and time.</small>
    </div>
    <section class="schedule" id="scheduleList" aria-label="Run schedule"></section>
  </main>

  <div class="modal-backdrop" id="modalBackdrop" aria-hidden="true">
    <article class="modal" role="dialog" aria-modal="true" aria-labelledby="modalTitle">
      <div class="modal-head">
        <div>
          <h3 id="modalTitle"></h3>
          <div class="modal-subtitle" id="modalSubtitle"></div>
          <div class="chipline" id="modalChips"></div>
        </div>
        <button class="modal-close" id="modalClose" type="button" aria-label="Close details">&times;</button>
      </div>
      <div class="modal-body" id="modalBody"></div>
    </article>
  </div>

  <script>
    const events = __EVENTS_JSON__;
    const categoryColors = __CATEGORY_COLORS_JSON__;
    const meta = __META_JSON__;
    const todayIso = meta.generatedOn;
    let currentMonth = meta.defaultMonth;

    const nameFilter = document.getElementById('nameFilter');
    const locationFilter = document.getElementById('locationFilter');
    const categoryFilter = document.getElementById('categoryFilter');
    const dateStart = document.getElementById('dateStart');
    const dateEnd = document.getElementById('dateEnd');
    const timeStart = document.getElementById('timeStart');
    const timeEnd = document.getElementById('timeEnd');
    const resetBtn = document.getElementById('resetBtn');
    const visibleCount = document.getElementById('visibleCount');
    const dateSpan = document.getElementById('dateSpan');
    const categoryCount = document.getElementById('categoryCount');
    const locationCount = document.getElementById('locationCount');
    const legend = document.getElementById('legend');
    const calendarGrid = document.getElementById('calendarGrid');
    const monthLabel = document.getElementById('monthLabel');
    const monthSelect = document.getElementById('monthSelect');
    const monthCount = document.getElementById('monthCount');
    const scheduleList = document.getElementById('scheduleList');
    const modalBackdrop = document.getElementById('modalBackdrop');
    const modalClose = document.getElementById('modalClose');

    function escapeHtml(value) {
      return String(value || '').replace(/[&<>"]/g, s => ({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;'}[s]));
    }

    function parseDate(iso) {
      return new Date(`${iso}T12:00:00`);
    }

    function dateToIso(date) {
      const year = date.getFullYear();
      const month = String(date.getMonth() + 1).padStart(2, '0');
      const day = String(date.getDate()).padStart(2, '0');
      return `${year}-${month}-${day}`;
    }

    function addDays(date, amount) {
      const copy = new Date(date);
      copy.setDate(copy.getDate() + amount);
      return copy;
    }

    function monthName(monthKey) {
      const [year, month] = monthKey.split('-').map(Number);
      return new Date(year, month - 1, 1).toLocaleDateString(undefined, { month: 'long', year: 'numeric' });
    }

    function formatDate(iso, options = { month: 'short', day: 'numeric', year: 'numeric' }) {
      return parseDate(iso).toLocaleDateString(undefined, options);
    }

    function dateRangeLabel(event) {
      if (event.startDate === event.endDate) return formatDate(event.startDate);
      return `${formatDate(event.startDate)} - ${formatDate(event.endDate)}`;
    }

    function timeToMinutes(value) {
      if (!value) return null;
      const [hours, minutes] = value.split(':').map(Number);
      return hours * 60 + minutes;
    }

    function timeLabel(value) {
      if (!value) return '';
      const [rawHour, rawMinute] = value.split(':').map(Number);
      const suffix = rawHour >= 12 ? 'PM' : 'AM';
      const hour = rawHour === 0 ? 12 : rawHour > 12 ? rawHour - 12 : rawHour;
      return `${hour}:${String(rawMinute).padStart(2, '0')} ${suffix}`;
    }

    function timeRange(event) {
      if (!event.startTime && !event.endTime) return 'All day';
      if (event.startTime && event.endTime) return `${timeLabel(event.startTime)} - ${timeLabel(event.endTime)}`;
      return timeLabel(event.startTime || event.endTime);
    }

    function hexToRgb(hex) {
      const clean = String(hex || '#64748B').replace('#', '');
      return {
        r: parseInt(clean.slice(0, 2), 16),
        g: parseInt(clean.slice(2, 4), 16),
        b: parseInt(clean.slice(4, 6), 16),
      };
    }

    function rgba(hex, alpha) {
      const c = hexToRgb(hex);
      return `rgba(${c.r}, ${c.g}, ${c.b}, ${alpha})`;
    }

    function sortedUnique(values) {
      return Array.from(new Set(values.filter(Boolean))).sort((a, b) => a.localeCompare(b));
    }

    function optionHtml(value, label) {
      return `<option value="${escapeHtml(value)}">${escapeHtml(label || value)}</option>`;
    }

    function populateControls() {
      const locations = sortedUnique(events.map(event => event.location));
      const categories = sortedUnique(events.map(event => event.category));
      locationFilter.innerHTML = optionHtml('all', 'All locations') + locations.map(value => optionHtml(value, value)).join('');
      categoryFilter.innerHTML = optionHtml('all', 'All categories') + categories.map(value => optionHtml(value, value)).join('');

      const months = sortedUnique(events.flatMap(event => {
        const keys = [];
        let cursor = new Date(parseDate(event.startDate).getFullYear(), parseDate(event.startDate).getMonth(), 1);
        const end = new Date(parseDate(event.endDate).getFullYear(), parseDate(event.endDate).getMonth(), 1);
        while (cursor <= end) {
          keys.push(dateToIso(cursor).slice(0, 7));
          cursor = new Date(cursor.getFullYear(), cursor.getMonth() + 1, 1);
        }
        return keys;
      }));
      monthSelect.innerHTML = months.map(month => optionHtml(month, monthName(month))).join('');
      monthSelect.value = currentMonth;
    }

    function filterState() {
      return {
        search: nameFilter.value.trim().toLowerCase(),
        location: locationFilter.value,
        category: categoryFilter.value,
        dateFrom: dateStart.value,
        dateTo: dateEnd.value,
        timeFrom: timeStart.value,
        timeTo: timeEnd.value,
      };
    }

    function overlapsDateRange(event, from, to) {
      const rangeStart = from || '0000-01-01';
      const rangeEnd = to || '9999-12-31';
      return event.startDate <= rangeEnd && event.endDate >= rangeStart;
    }

    function overlapsTimeRange(event, from, to) {
      if (!from && !to) return true;
      let start = timeToMinutes(event.startTime);
      let end = timeToMinutes(event.endTime);
      if (start === null && end === null) return true;
      if (start === null) start = 0;
      if (end === null) end = 1439;
      if (end < start) end += 1440;
      const rangeStart = from ? timeToMinutes(from) : 0;
      const rangeEnd = to ? timeToMinutes(to) : 1439;
      return start <= rangeEnd && end >= rangeStart;
    }

    function passes(event, filters) {
      const searchText = [
        event.title,
        event.location,
        event.category,
        event.type,
        event.status,
        event.team,
        event.eventId,
      ].join(' ').toLowerCase();
      const searchOk = !filters.search || searchText.includes(filters.search);
      const locationOk = filters.location === 'all' || event.location === filters.location;
      const categoryOk = filters.category === 'all' || event.category === filters.category;
      return searchOk
        && locationOk
        && categoryOk
        && overlapsDateRange(event, filters.dateFrom, filters.dateTo)
        && overlapsTimeRange(event, filters.timeFrom, filters.timeTo);
    }

    function filteredEvents() {
      const filters = filterState();
      return events.filter(event => passes(event, filters)).sort(compareEvents);
    }

    function compareEvents(a, b) {
      return `${a.startDate} ${a.startTime} ${a.title}`.localeCompare(`${b.startDate} ${b.startTime} ${b.title}`);
    }

    function renderSummary(filtered) {
      visibleCount.textContent = filtered.length;
      categoryCount.textContent = new Set(filtered.map(event => event.category)).size;
      locationCount.textContent = new Set(filtered.map(event => event.location)).size;
      if (!filtered.length) {
        dateSpan.textContent = '-';
        return;
      }
      const first = filtered.reduce((min, event) => event.startDate < min ? event.startDate : min, filtered[0].startDate);
      const last = filtered.reduce((max, event) => event.endDate > max ? event.endDate : max, filtered[0].endDate);
      dateSpan.textContent = first === last ? formatDate(first, { month: 'short', day: 'numeric', year: 'numeric' }) : `${formatDate(first, { month: 'short', day: 'numeric' })} - ${formatDate(last, { month: 'short', day: 'numeric', year: 'numeric' })}`;
    }

    function renderLegend(filtered) {
      const counts = filtered.reduce((map, event) => {
        map[event.category] = (map[event.category] || 0) + 1;
        return map;
      }, {});
      const categories = sortedUnique(Object.keys(categoryColors));
      legend.innerHTML = categories.map(category => `
        <span class="legend-item">
          <span class="swatch" style="--swatch:${categoryColors[category]}"></span>
          ${escapeHtml(category)} (${counts[category] || 0})
        </span>
      `).join('');
    }

    function renderCalendar(filtered) {
      monthLabel.textContent = monthName(currentMonth);
      monthSelect.value = currentMonth;
      const [year, month] = currentMonth.split('-').map(Number);
      const firstOfMonth = new Date(year, month - 1, 1);
      const start = addDays(firstOfMonth, -firstOfMonth.getDay());
      const monthStart = dateToIso(firstOfMonth);
      const monthEnd = dateToIso(new Date(year, month, 0));
      const monthEvents = filtered.filter(event => overlapsDateRange(event, monthStart, monthEnd));
      monthCount.textContent = `${monthEvents.length} event${monthEvents.length === 1 ? '' : 's'} in ${monthName(currentMonth)}`;

      const dayHeaders = ['Sun', 'Mon', 'Tue', 'Wed', 'Thu', 'Fri', 'Sat']
        .map(day => `<div class="month-dow">${day}</div>`).join('');
      const cells = [];
      for (let index = 0; index < 42; index += 1) {
        const day = addDays(start, index);
        const iso = dateToIso(day);
        const outside = day.getMonth() !== firstOfMonth.getMonth();
        const dayEvents = monthEvents.filter(event => event.startDate <= iso && event.endDate >= iso).sort(compareEvents);
        const visible = dayEvents.slice(0, 4);
        cells.push(`
          <div class="month-cell ${outside ? 'outside' : ''} ${iso === todayIso ? 'today' : ''}">
            <div class="month-date"><span>${day.getDate()}</span><span class="day-total">${dayEvents.length || ''}</span></div>
            ${visible.map(event => calendarEventHtml(event)).join('')}
            ${dayEvents.length > visible.length ? `<div class="month-more">+${dayEvents.length - visible.length} more</div>` : ''}
          </div>
        `);
      }
      calendarGrid.innerHTML = dayHeaders + cells.join('');
    }

    function calendarEventHtml(event) {
      const color = event.categoryColor;
      return `
        <button class="calendar-event" type="button" data-event-id="${event.id}" style="--event-color:${color};--event-bg:${rgba(color, .12)}">
          <strong>${escapeHtml(event.title)}</strong>
          <span>${escapeHtml(timeRange(event))}</span>
        </button>
      `;
    }

    function renderSchedule(filtered) {
      if (!filtered.length) {
        scheduleList.innerHTML = '<div class="empty">No events match the current filters.</div>';
        return;
      }
      scheduleList.innerHTML = filtered.map(event => `
        <article class="schedule-row" data-event-id="${event.id}" tabindex="0" role="button" style="--event-color:${event.categoryColor}">
          <div>
            <span class="date-pill">${escapeHtml(dateRangeLabel(event))}</span>
            <div class="timeblock">${escapeHtml(timeRange(event))}</div>
          </div>
          <div>
            <div class="titleline">${escapeHtml(event.title)}</div>
            <div class="details">${escapeHtml(event.type)}${event.team ? ` - ${escapeHtml(event.team)}` : ''}</div>
            <div class="chipline">
              <span class="chip">${escapeHtml(event.category)}</span>
              <span class="chip status-chip" style="--status-bg:${event.statusBackground};--status-fg:${event.statusText};">${escapeHtml(event.status)}</span>
              ${event.foundation === 'Yes' ? '<span class="chip">Foundation</span>' : ''}
            </div>
          </div>
          <div>
            <div class="modal-label">Location</div>
            <div class="location-text">${escapeHtml(event.location)}</div>
          </div>
          <div>
            <div class="modal-label">Category</div>
            <div class="category-pill"><span class="swatch" style="--swatch:${event.categoryColor}"></span>${escapeHtml(event.categoryCode || event.category)}</div>
          </div>
        </article>
      `).join('');
    }

    function render() {
      const filtered = filteredEvents();
      renderSummary(filtered);
      renderLegend(filtered);
      renderCalendar(filtered);
      renderSchedule(filtered);
    }

    function openModal(eventId) {
      const event = events.find(item => String(item.id) === String(eventId));
      if (!event) return;
      document.getElementById('modalTitle').textContent = event.title;
      document.getElementById('modalSubtitle').textContent = `${dateRangeLabel(event)} | ${timeRange(event)}`;
      document.getElementById('modalChips').innerHTML = `
        <span class="chip">${escapeHtml(event.category)}</span>
        <span class="chip status-chip" style="--status-bg:${event.statusBackground};--status-fg:${event.statusText};">${escapeHtml(event.status)}</span>
        <span class="chip">${escapeHtml(event.type)}</span>
      `;
      const fields = [
        ['Location', event.location],
        ['Event ID', event.eventId || 'Not listed'],
        ['Category Code', event.categoryCode || 'Not listed'],
        ['Type Code', event.typeCode || 'Not listed'],
        ['Team', event.team || 'Not listed'],
        ['Kallman Foundation Event', event.foundation],
        ['Last Updated', event.updated ? formatDate(event.updated) : 'Not listed'],
      ];
      document.getElementById('modalBody').innerHTML = fields.map(([label, value]) => `
        <div class="modal-field ${label === 'Team' ? 'wide' : ''}">
          <div class="modal-label">${escapeHtml(label)}</div>
          <div class="modal-value">${escapeHtml(value)}</div>
        </div>
      `).join('');
      modalBackdrop.classList.add('open');
      modalBackdrop.setAttribute('aria-hidden', 'false');
    }

    function closeModal() {
      modalBackdrop.classList.remove('open');
      modalBackdrop.setAttribute('aria-hidden', 'true');
    }

    function changeMonth(delta) {
      const [year, month] = currentMonth.split('-').map(Number);
      const next = new Date(year, month - 1 + delta, 1);
      currentMonth = dateToIso(next).slice(0, 7);
      render();
    }

    [nameFilter, locationFilter, categoryFilter, dateStart, dateEnd, timeStart, timeEnd].forEach(control => {
      control.addEventListener(control.tagName === 'INPUT' ? 'input' : 'change', () => {
        if (control === dateStart && dateStart.value) currentMonth = dateStart.value.slice(0, 7);
        render();
      });
    });
    monthSelect.addEventListener('change', () => {
      currentMonth = monthSelect.value;
      render();
    });
    document.getElementById('prevMonth').addEventListener('click', () => changeMonth(-1));
    document.getElementById('nextMonth').addEventListener('click', () => changeMonth(1));
    resetBtn.addEventListener('click', () => {
      nameFilter.value = '';
      locationFilter.value = 'all';
      categoryFilter.value = 'all';
      dateStart.value = '';
      dateEnd.value = '';
      timeStart.value = '';
      timeEnd.value = '';
      currentMonth = meta.defaultMonth;
      render();
    });
    document.addEventListener('click', event => {
      const target = event.target.closest('[data-event-id]');
      if (target) openModal(target.dataset.eventId);
      if (event.target === modalBackdrop) closeModal();
    });
    document.addEventListener('keydown', event => {
      if (event.key === 'Escape') closeModal();
      if (event.key === 'Enter') {
        const target = event.target.closest('[data-event-id]');
        if (target) openModal(target.dataset.eventId);
      }
    });
    modalClose.addEventListener('click', closeModal);

    populateControls();
    render();
  </script>
</body>
</html>
"""


def main() -> None:
    events, category_colors, _status_colors, metadata = build_events()
    serializable_metadata = dict(metadata)
    serializable_metadata["categoriesByCount"] = dict(metadata["categoriesByCount"])
    html = (
        HTML_TEMPLATE.replace("__EVENTS_JSON__", json.dumps(events, ensure_ascii=True))
        .replace("__CATEGORY_COLORS_JSON__", json.dumps(category_colors, ensure_ascii=True))
        .replace("__META_JSON__", json.dumps(serializable_metadata, ensure_ascii=True))
        .replace("__GENERATED_ON__", GENERATED_ON)
    )
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_PATH.write_text(html, encoding="utf-8")
    print(f"Wrote {OUTPUT_PATH}")
    print(f"Events: {len(events)}")
    print(f"Categories: {len(category_colors)}")
    print(f"Default month: {metadata['defaultMonth']}")


if __name__ == "__main__":
    main()
